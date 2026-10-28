# -*- coding: utf-8 -*-
"""XLSX (Excel) to HTML rendering with sheet tabs, merges and basic styles."""

from __future__ import annotations

import datetime as _dt
import html
import io
import logging
import math
from typing import Any

from .html_doc import xlsx_document

logger = logging.getLogger(__name__)

MAX_ROWS_PER_SHEET = 400
MAX_COLS_PER_SHEET = 60
MAX_SHEETS = 30
_COL_LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"


class XlsxRenderError(RuntimeError):
    """Raised when the workbook cannot be parsed."""


def _escape(value: Any) -> str:
    return html.escape(str(value) if value is not None else "", quote=False)


def _column_label(index: int) -> str:
    label = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        label = _COL_LETTERS[rem] + label
    return label


def _argb_to_hex(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    # Excel/openpyxl store ARGB. The alpha channel is not meaningful for a
    # preview (and openpyxl often writes "00" for opaque author colours), so
    # only the RGB part is used.
    if len(text) == 8:
        return f"#{text[2:]}"
    if len(text) == 6:
        return f"#{text}"
    return None


def _format_number(value: float, number_format: str) -> str:
    fmt = (number_format or "").lower()
    try:
        if "%" in fmt:
            decimals = 0
            if "." in fmt:
                decimals = min(6, len(fmt.split(".", 1)[1].split("%", 1)[0]))
            return f"{value * 100:.{decimals}f}%"
        currency = next(
            (symbol for symbol in ("¥", "$", "€", "£", "￥") if symbol in fmt),
            None,
        )
        if currency:
            decimals = 2
            if "." in fmt:
                decimals = min(6, len(fmt.split(".", 1)[1].split()[0]))
            return f"{currency}{value:,.{decimals}f}"
        if "0.00" in fmt:
            return f"{value:,.2f}"
        if "0.0" in fmt:
            return f"{value:,.1f}"
        if math.isfinite(value) and abs(value - round(value)) < 1e-9:
            return str(int(round(value)))
        return f"{value:.10g}"
    except (TypeError, ValueError, OverflowError):
        return str(value)


def _format_value(cell: Any) -> tuple[str, str]:
    """Return ``(display_text, css_classes)`` for one cell."""

    value = cell.value
    if value is None:
        return "", ""
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), "qop-bool"
    if isinstance(value, _dt.datetime):
        text = value.isoformat(sep=" ")
        if text.endswith(" 00:00:00"):
            text = text[:10]
        return text, "qop-num"
    if isinstance(value, _dt.date):
        return value.isoformat(), "qop-num"
    if isinstance(value, _dt.time):
        return value.isoformat(), "qop-num"
    if isinstance(value, (int, float)):
        return _format_number(float(value), getattr(cell, "number_format", "")), "qop-num"
    return str(value), ""


def _cell_style(cell: Any) -> str:
    declarations: list[str] = []
    try:
        font = cell.font
        if font is not None:
            if font.bold:
                declarations.append("font-weight:600")
            if font.italic:
                declarations.append("font-style:italic")
            if font.underline:
                declarations.append("text-decoration:underline")
            color = _argb_to_hex(getattr(font.color, "rgb", None))
            if color:
                declarations.append(f"color:{color}")
    except Exception:
        pass
    try:
        fill = cell.fill
        if fill is not None and getattr(fill, "patternType", None) == "solid":
            color = _argb_to_hex(getattr(fill.fgColor, "rgb", None))
            if color:
                declarations.append(f"background:{color}")
    except Exception:
        pass
    try:
        alignment = cell.alignment
        if alignment is not None:
            horizontal = (alignment.horizontal or "").lower()
            if horizontal in {"left", "center", "right", "justify"}:
                declarations.append(f"text-align:{horizontal}")
            if alignment.wrap_text:
                declarations.append("white-space:pre-wrap")
    except Exception:
        pass
    return ";".join(declarations)


def _merged_maps(worksheet: Any) -> tuple[dict[tuple[int, int], tuple[int, int]], set[tuple[int, int]]]:
    spans: dict[tuple[int, int], tuple[int, int]] = {}
    covered: set[tuple[int, int]] = set()
    try:
        ranges = list(worksheet.merged_cells.ranges)
    except Exception:
        return spans, covered
    for merged in ranges:
        try:
            min_row, max_row = merged.min_row, merged.max_row
            min_col, max_col = merged.min_col, merged.max_col
        except Exception:
            continue
        spans[(min_row, min_col)] = (
            max_row - min_row + 1,
            max_col - min_col + 1,
        )
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                if (row, col) != (min_row, min_col):
                    covered.add((row, col))
    return spans, covered


def _render_sheet(
    worksheet: Any,
    formula_sheet: Any,
    index: int,
    warnings: list[str],
) -> str | None:
    if getattr(worksheet, "sheet_state", "visible") != "visible":
        return None

    try:
        max_row = int(worksheet.max_row or 0)
        max_col = int(worksheet.max_column or 0)
    except Exception:
        max_row = max_col = 0
    if max_row <= 0 or max_col <= 0:
        return (
            f'<section class="qop-sheet"><h2>{_escape(worksheet.title)}</h2>'
            '<div class="qop-sheet-scroll"><p class="qop-empty" '
            'style="padding:16px">空工作表</p></div></section>'
        )

    rows_truncated = max_row > MAX_ROWS_PER_SHEET
    cols_truncated = max_col > MAX_COLS_PER_SHEET
    if rows_truncated:
        warnings.append(
            f"工作表「{worksheet.title}」共 {max_row} 行，仅显示前 "
            f"{MAX_ROWS_PER_SHEET} 行。",
        )
    if cols_truncated:
        warnings.append(
            f"工作表「{worksheet.title}」共 {max_col} 列，仅显示前 "
            f"{MAX_COLS_PER_SHEET} 列。",
        )
    row_count = min(max_row, MAX_ROWS_PER_SHEET)
    col_count = min(max_col, MAX_COLS_PER_SHEET)

    spans, covered = _merged_maps(worksheet)

    # Column widths (Excel character units -> approximate pixels).
    colgroup: list[str] = ['<col style="width:52px">']
    for col in range(1, col_count + 1):
        try:
            letter = _column_label(col - 1)
            width = worksheet.column_dimensions[letter].width
        except Exception:
            width = None
        px = 96 if not width else max(56, min(420, int(width * 7.6)))
        colgroup.append(f'<col style="width:{px}px">')

    header_cells = "".join(
        f"<th>{_column_label(col - 1)}</th>"
        for col in range(1, col_count + 1)
    )
    body_rows: list[str] = []
    for row in range(1, row_count + 1):
        cells: list[str] = [f'<th scope="row">{row}</th>']
        for col in range(1, col_count + 1):
            if (row, col) in covered:
                continue
            try:
                cell = worksheet.cell(row=row, column=col)
            except Exception:
                cells.append("<td></td>")
                continue
            text, classes = _format_value(cell)
            if not text and formula_sheet is not None:
                try:
                    formula_cell = formula_sheet.cell(row=row, column=col)
                    formula = formula_cell.value
                except Exception:
                    formula = None
                if isinstance(formula, str) and formula.startswith("="):
                    text = formula
                    classes = "qop-formula"
            attrs = ""
            span = spans.get((row, col))
            if span:
                rowspan, colspan = span
                if rowspan > 1:
                    attrs += f' rowspan="{rowspan}"'
                if colspan > 1:
                    attrs += f' colspan="{colspan}"'
            style = _cell_style(cell)
            if style:
                attrs += f' style="{style}"'
            class_attr = f' class="{classes}"' if classes else ""
            cells.append(f"<td{attrs}{class_attr}>{_escape(text)}</td>")
        body_rows.append(f"<tr>{''.join(cells)}</tr>")

    note = ""
    if rows_truncated or cols_truncated:
        note = (
            '<p class="qop-empty" style="margin:0;padding:8px 14px;'
            'border-top:1px solid var(--qop-border)">'
            "内容较长，已按上限截断以便快速预览。"
            "</p>"
        )
    return (
        f'<section class="qop-sheet" id="sheet-{index}">'
        f"<h2>{_escape(worksheet.title)}</h2>"
        f'<div class="qop-sheet-scroll"><table class="qop-grid">'
        f"<colgroup>{''.join(colgroup)}</colgroup>"
        f'<thead><tr><th class="qop-corner">#</th>{header_cells}</tr></thead>'
        f"<tbody>{''.join(body_rows)}</tbody></table></div>"
        f"{note}</section>"
    )


def render_xlsx(
    data: bytes,
    *,
    filename: str,
    theme: str = "auto",
) -> tuple[str, list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(
            io.BytesIO(data),
            data_only=True,
            read_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise XlsxRenderError(f"无法解析 XLSX 文件：{exc}") from exc

    formula_book = None
    # Reading the workbook twice is only worth it for reasonably sized files;
    # for huge workbooks the cached values are enough and memory matters more.
    if len(data) <= 12 * 1024 * 1024:
        try:
            formula_book = load_workbook(
                io.BytesIO(data),
                data_only=False,
                read_only=False,
                keep_links=False,
            )
        except Exception:
            formula_book = None

    warnings: list[str] = []
    sheets_html: list[str] = []
    try:
        sheet_names = list(workbook.sheetnames)
        if len(sheet_names) > MAX_SHEETS:
            warnings.append(
                f"工作簿共 {len(sheet_names)} 个工作表，仅显示前 {MAX_SHEETS} 个。",
            )
        hidden = 0
        for index, name in enumerate(sheet_names[:MAX_SHEETS], start=1):
            worksheet = workbook[name]
            if getattr(worksheet, "sheet_state", "visible") != "visible":
                hidden += 1
                continue
            formula_sheet = None
            if formula_book is not None:
                try:
                    formula_sheet = formula_book[name]
                except Exception:
                    formula_sheet = None
            try:
                rendered = _render_sheet(worksheet, formula_sheet, index, warnings)
            except Exception as exc:
                logger.info("Sheet %s failed to render: %s", name, exc)
                warnings.append(f"工作表「{name}」渲染失败，已跳过。")
                continue
            if rendered:
                sheets_html.append(rendered)
        if hidden:
            warnings.append(f"已跳过 {hidden} 个隐藏工作表。")
    finally:
        try:
            workbook.close()
        except Exception:
            pass
        if formula_book is not None:
            try:
                formula_book.close()
            except Exception:
                pass

    if not sheets_html:
        sheets_html.append(
            '<section class="qop-sheet"><h2>工作簿</h2>'
            '<p class="qop-empty" style="padding:16px">没有可显示的工作表。</p>'
            "</section>",
        )

    meta = (
        f'<span class="qop-chip">XLSX</span>'
        f'<span class="qop-chip">{_escape(filename)}</span>'
        f'<span class="qop-chip">{len(sheet_names)} 个工作表</span>'
    )
    doc = xlsx_document(
        f'<div class="qop-workbook">{"".join(sheets_html)}</div>',
        title=filename,
        meta=meta,
        warnings=" ".join(_escape(item) for item in warnings),
        theme=theme,
    )
    return doc, warnings
