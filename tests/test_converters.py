# -*- coding: utf-8 -*-
"""Converter unit tests (fixtures are generated with the same libraries)."""

from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches

from converter import convert_office, detect_kind
from converter.service import (
    FileTooLargeError,
    OfficePreviewError,
    UnsupportedFormatError,
    clear_cache,
)


@pytest.fixture(autouse=True)
def _clear_cache():
    clear_cache()
    yield
    clear_cache()


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("季度报告", 0)
    doc.add_heading("一、总体情况", level=1)
    paragraph = doc.add_paragraph("本季度围绕")
    paragraph.add_run("重点任务").bold = True
    paragraph.add_run("推进工作。")
    doc.add_paragraph("第一项", style="List Bullet")
    doc.add_paragraph("第二项", style="List Bullet")
    doc.add_paragraph("第一步", style="List Number")
    doc.add_paragraph("第二步", style="List Number")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "收入"
    table.cell(1, 1).text = "100"
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["地区", "金额", "占比"])
    ws.append(["华东", 12345.67, 0.35])
    ws.append(["华北", 2345, 0.12])
    ws["C2"].number_format = "0.0%"
    ws["C3"].number_format = "0.0%"
    ws["B2"].number_format = "¥#,##0.00"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF2563EB")
    ws.merge_cells("A5:C5")
    ws["A5"] = "合并标题"
    wb.create_sheet("说明")["A1"] = "第二张表"
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def make_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "年度规划"
    slide.placeholders[1].text = "第一部分：市场\n第二部分：产品"
    table_slide = prs.slides.add_slide(prs.slide_layouts[5])
    table_slide.shapes.title.text = "数据"
    table = table_slide.shapes.add_table(
        2,
        2,
        Inches(1),
        Inches(2),
        Inches(5),
        Inches(1.5),
    ).table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "Q1"
    table.cell(1, 0).text = "收入"
    table.cell(1, 1).text = "100"
    table_slide.notes_slide.notes_text_frame.text = "备注内容"
    stream = io.BytesIO()
    prs.save(stream)
    return stream.getvalue()


def test_docx_renders_semantic_html():
    result = convert_office("报告.docx", make_docx())
    html = result["html"]
    assert result["kind"] == "docx"
    assert result["format"] == "html"
    assert "季度报告" in html
    assert "重点任务" in html
    assert "<table" in html
    assert "<ul>" in html and "<ol>" in html
    assert html.startswith("<!doctype html>")
    assert "qop-doc" in html


def test_xlsx_renders_sheets_formats_and_merges():
    result = convert_office("报表.xlsx", make_xlsx())
    html = result["html"]
    assert result["kind"] == "xlsx"
    for token in [
        "数据",
        "说明",
        "华北",
        "35.0%",
        "¥12,345.67",
        'colspan="3"',
        "qop-grid",
    ]:
        assert token in html, token


def test_pptx_renders_slides_tables_and_notes():
    result = convert_office("演示.pptx", make_pptx())
    html = result["html"]
    assert result["kind"] == "pptx"
    for token in [
        "年度规划",
        "第一部分",
        "qop-slide",
        "幻灯片导航",
        "备注内容",
        "<table",
    ]:
        assert token in html, token


def test_csv_rendering_and_delimiter_sniffing():
    result = convert_office("people.csv", "姓名;年龄\n张三;30\n".encode())
    html = result["html"]
    assert result["kind"] == "csv"
    assert "姓名" in html and "张三" in html
    assert "qop-grid" in html


def test_theme_is_forced_when_requested():
    data = make_docx()
    light = convert_office("a.docx", data, theme="light")["html"]
    dark = convert_office("a.docx", data, theme="dark")["html"]
    assert 'class="qop-theme-light"' in light
    assert 'class="qop-theme-dark"' in dark


def test_cache_returns_cached_result():
    data = make_docx()
    first = convert_office("a.docx", data)
    second = convert_office("a.docx", data)
    assert first["cached"] is False
    assert second["cached"] is True
    assert second["html"] == first["html"]


def test_mislabelled_docx_is_detected_by_magic():
    kind, suffix = detect_kind("old.doc", make_docx())
    assert kind == "docx"
    assert suffix == ".doc"
    # Rendering must not require LibreOffice for a mislabelled modern file.
    result = convert_office("old.doc", make_docx())
    assert result["kind"] == "docx"


def test_legacy_ole_file_needs_libreoffice(monkeypatch):
    monkeypatch.setattr("converter.service.find_libreoffice", lambda: None)
    ole_header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512
    with pytest.raises(OfficePreviewError) as info:
        convert_office("old.doc", ole_header)
    assert info.value.code == "legacy_converter_missing"
    assert "LibreOffice" in str(info.value)


def test_zip_bomb_is_rejected(monkeypatch):
    monkeypatch.setattr("converter.service.MAX_UNCOMPRESSED_BYTES", 1024)
    with pytest.raises(OfficePreviewError) as info:
        convert_office("bomb.docx", make_docx())
    assert info.value.code == "archive_too_large"


def test_unsupported_extension_is_rejected():
    with pytest.raises(UnsupportedFormatError):
        convert_office("photo.png", b"\x89PNG\r\n")


def test_corrupted_document_reports_error():
    with pytest.raises(OfficePreviewError):
        convert_office("broken.docx", b"PK\x03\x04not really a zip")


def test_size_limit(monkeypatch):
    monkeypatch.setenv("QWENPAW_OFFICE_PREVIEW_MAX_MB", "1")
    with pytest.raises(FileTooLargeError):
        convert_office("big.docx", b"0" * (1024 * 1024 + 1))


def test_empty_file_rejected():
    with pytest.raises(OfficePreviewError):
        convert_office("empty.docx", b"")


def test_xlsx_formula_fallback_when_no_cached_value():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A2"] = "=A1+1"
    stream = io.BytesIO()
    wb.save(stream)
    html = convert_office("formula.xlsx", stream.getvalue())["html"]
    assert "=A1+1" in html


def test_docx_dates_and_unicode():
    doc = Document()
    doc.add_paragraph("中文 · emoji 📄 · <script>alert(1)</script>")
    stream = io.BytesIO()
    doc.save(stream)
    html = convert_office("x.docx", stream.getvalue())["html"]
    assert "中文" in html
    assert "<script>alert" not in html


def test_pptx_empty_deck_is_renderable():
    prs = Presentation()
    stream = io.BytesIO()
    prs.save(stream)
    result = convert_office("empty.pptx", stream.getvalue())
    assert result["kind"] == "pptx"
    assert "qop-slides" in result["html"]


def test_xlsx_date_values_are_text():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = dt.date(2026, 3, 1)
    ws["A2"] = dt.datetime(2026, 3, 2, 10, 30)
    stream = io.BytesIO()
    wb.save(stream)
    html = convert_office("dates.xlsx", stream.getvalue())["html"]
    assert "2026-03-01" in html
    assert "2026-03-02" in html


FIXTURES = Path(__file__).parent / "fixtures"


def test_committed_fixtures_cover_all_modern_formats():
    cases = {
        "sample.docx": ["Offline Sample Report", "offline smoke tests", "<table"],
        "sample.xlsx": ["Data", "Notes", "35.0%", "¥12,345.67", "Merged Title"],
        "sample.pptx": [
            "Offline Sample Deck",
            "First part",
            "Table Slide",
            "Speaker note for offline test",
            "data:image/",
        ],
        "sample.csv": ["Alice", "90"],
    }
    for name, markers in cases.items():
        result = convert_office(name, (FIXTURES / name).read_bytes())
        for marker in markers:
            assert marker in result["html"], (name, marker)


def test_native_pptx_renderer_reads_layout_inheritance():
    from converter.pptx_native import render_pptx_native

    html, warnings = render_pptx_native(
        (FIXTURES / "sample.pptx").read_bytes(),
        filename="sample.pptx",
        theme="light",
    )
    assert warnings == []
    # Placeholder positions come from slideLayout/slideMaster, not the slide.
    assert "Offline Sample Deck" in html
    assert "First part" in html
    assert "qop-shape" in html
    assert "Speaker note for offline test" in html
    assert "data:image/" in html


def test_render_pptx_falls_back_to_native(monkeypatch):
    import converter.pptx as pptx_module

    def boom(*_args, **_kwargs):
        raise ImportError("python-pptx is not installed")

    monkeypatch.setattr(pptx_module, "_render_with_python_pptx", boom)
    html, warnings = pptx_module.render_pptx(
        (FIXTURES / "sample.pptx").read_bytes(),
        filename="sample.pptx",
    )
    assert "Offline Sample Deck" in html
    assert any("python-pptx" in item for item in warnings)
